# -*- coding: utf-8 -*-
"""自审：越界、溢出估算、图片数、备注完整性"""
import math
from pptx import Presentation
from pptx.util import Emu

P = r"D:\isac\paper\ppt_out\final_presentation_cn.pptx"
prs = Presentation(P)
SW = prs.slide_width
SH = prs.slide_height

def emu2in(v):
    return round(v / 914400, 2)

issues = []
print(f"slides={len(prs.slides)}  size={emu2in(SW)}x{emu2in(SH)} in")
for i, s in enumerate(prs.slides, 1):
    pics = 0
    for sh in s.shapes:
        l, t, w, h = sh.left, sh.top, sh.width, sh.height
        if l is None:
            continue
        if l < 0 or t < 0 or l + w > SW + 1000 or t + h > SH + 1000:
            issues.append(f"S{i}: 越界 {sh.shape_type} at ({emu2in(l)},{emu2in(t)}) size ({emu2in(w)}x{emu2in(h)})")
        if sh.shape_type == 13:
            pics += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            # 估算所需高度
            need = 0.0
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    continue
                size = max((r.font.size.pt if r.font.size else 14) for r in p.runs)
                # 中文字宽约 1.0em，英文约 0.55em
                width_pt = emu2in(w) * 72
                em = sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in txt)
                per_line = max(1.0, width_pt / (size * 1.0))
                lines = max(1, math.ceil(em / per_line))
                need += lines * size * 1.32
            need_in = need / 72 + 0.05
            if need_in > emu2in(h) + 0.12:
                issues.append(f"S{i}: 可能溢出 需{need_in:.2f}in > 框{emu2in(h):.2f}in  «{tf.text[:28]}»")
    has_notes = s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() != ""
    print(f"  S{i:02d} 图片={pics}  备注={'Y' if has_notes else 'N'}")

print("\n=== 发现问题 ===")
if not issues:
    print("无")
for x in issues:
    print(" -", x)
