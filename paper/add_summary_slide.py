# -*- coding: utf-8 -*-
"""在 S17 之后插入一页：核心工作与阶段性结论"""
import copy
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = Path(r"C:\Users\Administrator\.dsh\attachments\v1\files\11\11e24751c46c01c933c651c3b061cb7f738b99c3aef10715a76d1034125f816f\开题ppt_叶嘉锐.pptx")
DST = Path(r"D:\isac\paper\ppt_out\开题ppt_叶嘉锐_加页版.pptx")
DST.parent.mkdir(parents=True, exist_ok=True)

INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
GREY_L = RGBColor(0xD9, 0xD9, 0xD9)
BLUE_D = RGBColor(0x1F, 0x4E, 0x79)
BLUE_L = RGBColor(0xDC, 0xE6, 0xEF)
WARM_D = RGBColor(0xB5, 0x53, 0x3C)
WARM_L = RGBColor(0xF6, 0xE8, 0xE4)
GREEN_D = RGBColor(0x4C, 0x7A, 0x5A)
GREEN_L = RGBColor(0xE4, 0xEE, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_CN = "微软雅黑"
FONT_EN = "Arial"

SW, SH = 13.333, 7.5
ML = 0.85
CW = SW - 2 * ML

prs = Presentation(str(SRC))


def set_font(run, size=14, bold=False, color=INK, ea=FONT_CN, latin=FONT_EN):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", ea), ("a:cs", ea)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", face)


def textbox(slide, left, top, width, height, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    return tf


def put(tf, text, size=14, bold=False, color=INK, first=False, line=1.25,
        space_before=0, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    set_font(r, size, bold, color)
    return p


def rect(slide, left, top, width, height, fill=None, line=None, lw=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def rrect(slide, left, top, width, height, fill=None, line=None, lw=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    sh.adjustments[0] = 0.06
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


blank = prs.slide_layouts[6]
s = prs.slides.add_slide(blank)

# ---------- 统一的页眉样式（对齐原 deck）----------
tf = textbox(s, ML, 0.30, CW, 0.28)
put(tf, "上海海事大学", 10, False, GREY, first=True)
tf = textbox(s, ML, 0.55, CW, 0.28)
put(tf, "面向实时视频业务的 OFDM-ISAC｜跨层传输与感知共存研究", 10, False, GREY, first=True)

tf = textbox(s, ML, 0.95, CW, 0.55)
put(tf, "核心工作与阶段性结论", 26, True, INK, first=True)
rect(s, ML, 1.60, 1.1, 0.045, fill=BLUE_D)
tf = textbox(s, ML + 1.3, 1.52, CW - 1.3, 0.30)
put(tf, "从「跑通」到「定位」：四件可以写进论文的事", 12, False, GREY, first=True)

# ---------- 四栏 ----------
cards = [
    ("01", "实时视频业务试验床", BLUE_D, BLUE_L, [
        "摄像头采集 → H.264 → 裸流 → ISAC",
        "与平台示例（文件推流）三处不同",
        "端到端时延测量与逐层分解",
    ]),
    ("02", "分帧元数据的单点故障", WARM_D, WARM_L, [
        "2 字节长度前缀位于 RS 码字内部",
        "前缀裸露时单比特错误即整流失步",
        "外码稳态功能是分帧，非纠错",
    ]),
    ("03", "容器与解码器的放大", BLUE_D, BLUE_L, [
        "同一误码流，四栈对比实测",
        "MPEG-TS+VLC ≈500× vs 裸流+GStreamer ≈4×",
        "γ = γ_header × γ_state 双因子",
    ]),
    ("04", "缓冲策略与时延", GREEN_D, GREEN_L, [
        "四级队列 32 → 4–8 帧",
        "阻塞写改为满则丢弃",
        "250 ms 量级，无累积趋势",
    ]),
]

gap = 0.22
w = (CW - 3 * gap) / 4
x = ML
top = 2.00
h = 3.55

for tag, title, cd, cl, bullets in cards:
    rrect(s, x, top, w, h, fill=cl, line=None)
    tf = textbox(s, x + 0.22, top + 0.20, w - 0.44, 0.30)
    put(tf, tag, 13, True, cd, first=True)
    tf = textbox(s, x + 0.22, top + 0.52, w - 0.44, 0.75)
    put(tf, title, 14.5, True, INK, first=True, line=1.2)
    rect(s, x + 0.22, top + 1.34, 0.5, 0.03, fill=cd)
    yy = top + 1.52
    for b in bullets:
        tf = textbox(s, x + 0.22, yy, 0.14, 0.22)
        put(tf, "·", 14, True, cd, first=True)
        tf = textbox(s, x + 0.40, yy, w - 0.62, 0.62)
        put(tf, b, 10.5, False, INK, first=True, line=1.30)
        yy += 0.64
    x += w + gap

# ---------- 底部：方法学一句 ----------
rect(s, ML, 5.78, CW, 0.012, fill=GREY_L)
tf = textbox(s, ML, 5.92, CW, 0.42)
put(tf, "方法学：测量链本身也需要被校验——外码实现的三处缺陷正是在自检中被定位，"
        "相关结论随之撤回并进入重测。", 11.5, False, GREY, first=True, line=1.30)

tf = textbox(s, ML, 6.45, CW, 0.40)
put(tf, "四件事都不依赖感知算法的开发，均为通信侧的系统级实测。", 13, True, BLUE_D, first=True)

# ---------- 页码位置对齐（右下）----------
tf = textbox(s, SW - ML - 1.0, 7.05, 1.0, 0.26, align=PP_ALIGN.RIGHT)
put(tf, "18", 9, False, GREY, first=True, align=PP_ALIGN.RIGHT)

# ---------- 移动到第 18 位（S17 之后）----------
xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)
xml_slides.remove(slides[-1])
xml_slides.insert(17, slides[-1])

prs.save(str(DST))
print("saved:", DST, "slides:", len(prs.slides))
