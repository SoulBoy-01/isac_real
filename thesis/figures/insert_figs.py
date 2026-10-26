# -*- coding: utf-8 -*-
"""把三张架构图插入开题报告，并顺延图号：
   图1 系统拓扑（替换原平面框图）
   图2 级联纠错编码链路架构（新增，置于研究内容 1）
   图3 容器放大因子 γ（原图 2）
   图4 微动感知处理架构（新增，置于研究内容 3）
   图5 技术路线总图（原图 3）
   图6 端到端时延实测（原图 4）
"""
import sys, io, copy, struct
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC = r"D:\isac\thesis\开题报告_v2_答辩版.docx"
DST = r"D:\isac\thesis\开题报告_v3_图文版.docx"
FIG = r"D:\isac\thesis\figures"

doc = Document(SRC)

# 可用正文宽度
sec = doc.sections[0]
usable = sec.page_width - sec.left_margin - sec.right_margin
maxw = min(int(usable * 0.98), int(Inches(6.3)))
print("可用正文宽度:", round(usable / 914400, 2), "in -> 图片宽度用", round(maxw / 914400, 2), "in")


def locate(needle):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise SystemExit(f"!!! 未找到: {needle}")


def replace_sub(text):
    """对单段文字做图号顺延（顺序敏感）"""
    rules = [("见图 4", "见图 6"), ("见图 3", "见图 5"),
             ("图 4　", "图 6　"), ("图 3　", "图 5　"), ("图 2　", "图 3　")]
    for a, b in rules:
        text = text.replace(a, b)
    return text


def png_ratio(path):
    """从 PNG 头部读取真实宽高，返回 高/宽，避免硬编码比例失真。"""
    with open(path, "rb") as f:
        head = f.read(26)
    w, h = struct.unpack(">II", head[16:24])
    return h / w


# ── 1) 图号顺延（先做，避免与新增段落混淆）──
for p in doc.paragraphs:
    if "图" not in p.text:
        continue
    old = p.text
    new = replace_sub(old)
    if new == old:
        continue
    # 尽量在 run 内替换；失败则压平到首 run
    done = False
    for r in p.runs:
        if any(a in r.text for a in ("见图 4", "见图 3", "图 4　", "图 3　", "图 2　")):
            r.text = replace_sub(r.text)
            done = True
    if not done:
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    print("  [renumber]", old.strip()[:42], "->", p.text.strip()[:42])

# ── 2) 替换图 1 的图片内容（保持 rId，仅换 blob 与显示尺寸）──
topo = f"{FIG}\\fig_topology_v2.png"
with open(topo, "rb") as f:
    topo_bytes = f.read()

img_part = None
for rid, part in doc.part.related_parts.items():
    if str(part.partname).endswith("image2.png"):
        img_part = part
        target_rid = rid
        break
assert img_part is not None, "未找到 image2.png 部件"
img_part._blob = topo_bytes
print("  [replace] image2.png ->", len(topo_bytes) // 1024, "KB")

for shp in doc.inline_shapes:
    blip = shp._inline.graphic.graphicData.pic.blipFill.blip
    if blip.get(qn("r:embed")) == target_rid:
        ratio = png_ratio(topo)
        shp.width = Emu(maxw)
        shp.height = Emu(int(maxw * ratio))
        print("  [resize] 图1 ->", round(maxw / 914400, 2), "x", round(maxw * ratio / 914400, 2), "in")
        break


def add_figure(anchor_p, img_path, caption, ratio):
    """在 anchor_p 之后插入：图片段 + 题注段"""
    new_p = copy.deepcopy(anchor_p._p)
    anchor_p._p.addnext(new_p)
    imgp = Paragraph(new_p, anchor_p._parent)
    for r in list(imgp.runs):
        r._r.getparent().remove(r._r)
    imgp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imgp.add_run().add_picture(img_path, width=Emu(maxw), height=Emu(int(maxw * ratio)))

    cap_tpl = locate("图 1　OpenISAC")          # 用现有题注段作模板（居中）
    cap_p = copy.deepcopy(cap_tpl._p)
    new_p.addnext(cap_p)
    capp = Paragraph(cap_p, cap_tpl._parent)
    for r in list(capp.runs):
        r._r.getparent().remove(r._r)
    capp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = capp.add_run(caption)
    ref = cap_tpl.runs[0] if cap_tpl.runs else None
    if ref is not None:
        run.bold = ref.bold
        run.font.size = ref.font.size
        run.font.name = ref.font.name
    print("  [figure]", img_path.split("\\")[-1], "| 题注:", caption[:30])


# ── 3) 图 2：级联纠错编码链路（置于研究内容 1，接收侧译码链之后）──
anchor = locate("需特别指出：内码译码器既不返回收敛标志")
anchor.runs[-1].text = anchor.runs[-1].text + \
    "级联编解码结构与残差判据的统计位置如图 2 所示。"
add_figure(anchor, f"{FIG}\\fig_coding_chain.png",
           "图 2　级联纠错编码链路架构：RS(255, K) 外码与 LDPC(1008, 504) 内码的级联编解码流程，"
           "虚线框内为残差判据点（RS 伴随式）",
           png_ratio(f"{FIG}\\fig_coding_chain.png"))

# ── 4) 图 4：微动感知处理架构（置于研究内容 3）──
anchor2 = locate("拟采用方案：利用接收链路已有的导频相位跟踪")
anchor2.runs[-1].text = anchor2.runs[-1].text + "处理流程与相位灵敏度预算如图 4 所示。"
add_figure(anchor2, f"{FIG}\\fig_sensing.png",
           "图 4　非接触微动感知处理架构：由目标距离门相位到呼吸/心跳参数估计的处理流程与灵敏度预算",
           png_ratio(f"{FIG}\\fig_sensing.png"))

doc.save(DST)
print("\nSAVED ->", DST)

# ── 5) 统一所有插图宽度到可用正文宽度（保持各自宽高比）──
d2 = Document(DST)
for s in d2.inline_shapes:
    ratio = s.height / s.width
    s.width = Emu(maxw)
    s.height = Emu(int(maxw * ratio))
d2.save(DST)
print("已统一插图宽度：")
for i, s in enumerate(Document(DST).inline_shapes):
    print(f"  img{i + 1}: {round(s.width / 914400, 2)} x {round(s.height / 914400, 2)} in")
