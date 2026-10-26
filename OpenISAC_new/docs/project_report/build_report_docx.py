# -*- coding: utf-8 -*-
"""Build the OpenISAC project acceptance report (Word) from the markdown body."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPORT_DIR = Path(__file__).resolve().parent
BODY_MD = REPORT_DIR / "OpenISAC_项目验收报告_body.md"
OUT_DOCX = REPORT_DIR / "OpenISAC_项目验收报告.docx"

FALLBACK_FONT = "Times New Roman"
CJK_FONT = "SimSun"
HEI_FONT = "SimHei"


def set_run_font(run, east_asia=CJK_FONT, latin=FALLBACK_FONT, size=None, bold=None, color=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), east_asia)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)


def add_heading(doc, text, level):
    h = doc.add_heading("", level=level)
    sizes = {0: 22, 1: 18, 2: 15, 3: 13}
    run = h.add_run(text)
    color = (0x0F, 0x3B, 0x66) if level >= 2 else (0x08, 0x2A, 0x4A)
    set_run_font(run, east_asia=HEI_FONT, latin="Arial", size=sizes.get(level, 12),
                 bold=True, color=color)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return h


def add_para(doc, text, size=10.5, bold=False, align=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    run.font.italic = italic
    return p


def add_md_runs(p, md_text, size=10.5):
    """Render a limited markdown subset (bold **...**, italic *...*) inside a paragraph."""
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', md_text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            run = p.add_run(tok[2:-2])
            set_run_font(run, size=size, bold=True)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            run = p.add_run(tok[1:-1])
            set_run_font(run, size=size)
            run.font.italic = True
        else:
            run = p.add_run(tok)
            set_run_font(run, size=size)


def add_md_paragraph(doc, md_text, size=10.5, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(0.74)
    add_md_runs(p, md_text, size=size)
    return p


def add_table(doc, rows, widths=None, header=True, font_size=9.5):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(cell_text.strip())
            is_header = header and i == 0
            set_run_font(run, east_asia=(HEI_FONT if is_header else CJK_FONT),
                         latin=FALLBACK_FONT, size=font_size, bold=is_header)
            if is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if widths and j < len(widths):
                cell.width = widths[j]
    return table


def add_image_caption(doc, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(caption_text)
    set_run_font(run, east_asia=HEI_FONT, size=9, bold=True, color=(0x44, 0x44, 0x44))
    return p


def add_image(doc, image_name, caption, width_cm=15.5):
    path = REPORT_DIR / image_name
    if not path.exists():
        print(f"[warn] missing image: {path}", file=sys.stderr)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_image_caption(doc, caption)


def add_eq(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, east_asia=CJK_FONT, latin="Cambria Math", size=10.5)
    run.font.italic = True
    return p


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run()
    fldChar1 = run._element.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
    instrText = run._element.makeelement(qn('w:instrText'), {})
    instrText.text = 'PAGE'
    fldChar2 = run._element.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'end'})
    run._element.append(fldChar1)
    run._element.append(instrText)
    run._element.append(fldChar2)
    set_run_font(run, size=9)


def add_toc(doc):
    h = doc.add_paragraph()
    run = h.add_run("目  录")
    set_run_font(run, east_asia=HEI_FONT, latin="Arial", size=16, bold=True,
                 color=(0x08, 0x2A, 0x4A))
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run()
    fld_begin = r._element.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
    instr = r._element.makeelement(qn('w:instrText'), {})
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = r._element.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'separate'})
    hint = r._element.makeelement(qn('w:t'), {})
    hint.text = "（打开后请右键此处选择“更新域”以生成目录）"
    fld_end = r._element.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'end'})
    r._element.append(fld_begin)
    r._element.append(instr)
    r._element.append(fld_sep)
    r._element.append(hint)
    r._element.append(fld_end)
    set_run_font(r, size=10)


# ----------------------------------------------------------------------
# Cover page
# ----------------------------------------------------------------------

def build_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    add_para(doc, "面向 6G 的实时 OFDM 通感一体化（ISAC）", size=26, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "开源实验平台研发项目", size=26, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
    add_para(doc, "项目验收报告", size=34, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)
    add_para(doc, "项目名称：OpenISAC 实时 OFDM 通感一体化实验平台",
             size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_para(doc, "代码仓库：https://github.com/SoulBoy-01/OpenISAC_new",
             size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_para(doc, "报告日期：2026 年 10 月", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=60)
    add_para(doc, "（提交企业项目负责人审阅）", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    doc.add_page_break()


# ----------------------------------------------------------------------
# Body builder: walks the markdown-ish body file
# ----------------------------------------------------------------------

def iter_body_blocks(md_text):
    lines = md_text.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        if line.startswith('#### '):
            yield ('h4', line[5:].strip())
        elif line.startswith('### '):
            yield ('h3', line[4:].strip())
        elif line.startswith('## '):
            yield ('h2', line[3:].strip())
        elif line.startswith('# '):
            yield ('h1', line[2:].strip())
        elif line.startswith('<!--PAGEBREAK-->'):
            yield ('pagebreak', None)
        elif line.startswith('<!--IMAGE:'):
            # <!--IMAGE:file.png|caption|width-->
            m = re.match(r'<!--IMAGE:([^|]+)\|([^|]*)\|?([0-9.]*)-->', line)
            if m:
                yield ('image', (m.group(1), m.group(2), float(m.group(3) or 15.5)))
        elif line.startswith('|'):
            tbl = []
            while i < n and lines[i].startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in row):
                    tbl.append(row)
                i += 1
            yield ('table', tbl)
            continue
        elif line.strip() == '':
            pass
        elif line.lstrip().startswith('$ '):
            yield ('eq', line.lstrip()[2:].strip())
        elif line.startswith('> '):
            yield ('quote', line[2:].strip())
        elif line.lstrip().startswith('- '):
            yield ('bullet', line.lstrip()[2:].strip())
        else:
            yield ('para', line.strip())
        i += 1


def build_body(doc, md_text):
    for kind, payload in iter_body_blocks(md_text):
        if kind == 'h1':
            add_heading(doc, payload, level=1)
        elif kind == 'h2':
            add_heading(doc, payload, level=2)
        elif kind == 'h3':
            add_heading(doc, payload, level=3)
        elif kind == 'h4':
            p = add_para(doc, payload, size=11, bold=True, space_after=3)
        elif kind == 'para':
            add_md_paragraph(doc, payload)
        elif kind == 'bullet':
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            add_md_runs(p, payload, size=10.5)
        elif kind == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(payload)
            set_run_font(run, size=10)
            run.font.italic = True
        elif kind == 'table':
            rows = payload
            if rows:
                add_table(doc, rows)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == 'eq':
            add_eq(doc, payload)
        elif kind == 'image':
            name, caption, width = payload
            add_image(doc, name, caption, width_cm=width)
        elif kind == 'pagebreak':
            doc.add_page_break()


def main():
    doc = Document()
    # base style
    style = doc.styles['Normal']
    style.font.name = FALLBACK_FONT
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), CJK_FONT)

    for section in doc.sections:
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.4)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.6)

    add_page_number_footer(doc)
    build_cover(doc)
    add_toc(doc)
    doc.add_page_break()

    md_text = BODY_MD.read_text(encoding='utf-8')
    build_body(doc, md_text)

    doc.save(str(OUT_DOCX))
    print(f"saved: {OUT_DOCX}")


if __name__ == '__main__':
    main()
